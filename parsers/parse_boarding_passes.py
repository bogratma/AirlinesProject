from __future__ import annotations

import re
import tempfile
import zipfile
from concurrent.futures import ProcessPoolExecutor
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

CELLS = {
    "sequence": "H1",
    "title": "A3",
    "passenger_name_raw": "B3",
    "ff_raw": "F3",
    "travel_class": "H3",
    "flight_no": "A5",
    "from_city": "D5",
    "to_city": "H5",
    "gate": "B7",
    "from_iata": "D7",
    "to_iata": "H7",
    "flight_date": "A9",
    "flight_time": "C9",
    "operator_note": "E9",
    "seat": "H11",
    "pnr": "B13",
    "ticket_no": "E13",
}

MARKERS = {
    "A1": "BOARDING PASS",
    "F1": "SEQUENCE:",
    "A7": "GATE",
    "G11": "SEAT",
    "A13": "PNR:",
    "D13": "E-TICKET",
}

CONSTANT_CANDIDATES = ["gate", "seat", "operator_note"]

GENDER = {"MR": "M", "MRS": "F", "MS": "F", "MISS": "F"}



def _get(rows, addr):
    col, row = coordinate_from_string(addr)
    try:
        v = rows[row - 1][column_index_from_string(col) - 1]
    except IndexError:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _as_date_str(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return None if v is None else str(v).strip()


def _as_time_str(v):
    if isinstance(v, (datetime, time)):
        return v.strftime("%H:%M")
    return None if v is None else str(v).strip()


def _as_digits(v, width):
    if v is None:
        return None
    if isinstance(v, float):
        v = int(v)
    s = str(v).strip()
    return s.zfill(width) if s.isdigit() else s


def parse_sheet(ws):
    rows = list(ws.iter_rows(min_row=1, max_row=13, max_col=8, values_only=True))

    for addr, expected in MARKERS.items():
        got = _get(rows, addr)
        if got is None or str(got).strip().upper() != expected:
            raise ValueError(f"layout mismatch at {addr}: expected {expected!r}, got {got!r}")

    rec = {field: _get(rows, addr) for field, addr in CELLS.items()}
    rec["flight_date"] = _as_date_str(rec["flight_date"])
    rec["flight_time"] = _as_time_str(rec["flight_time"])
    rec["ticket_no"] = _as_digits(rec["ticket_no"], 16)
    return rec


def parse_workbook(path: str):
    records, errors = [], []
    fname = Path(path).name
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # битый/запароленный файл
        return [], [{"file": fname, "sheet": None, "error": f"cannot open: {e}"}]
    try:
        for idx, ws in enumerate(wb.worksheets, start=1):
            try:
                rec = parse_sheet(ws)
                rec["source_file"] = fname
                rec["sheet"] = ws.title
                rec["sheet_idx"] = idx
                records.append(rec)
            except Exception as e:
                errors.append({"file": fname, "sheet": ws.title, "error": str(e)})
    finally:
        wb.close()
    return records, errors



def _name_key(raw):
    if not isinstance(raw, str):
        return None
    s = re.sub(r"[’'`]", "", raw.upper())
    toks = [t for t in s.split() if len(t) > 1]
    return " ".join(sorted(toks)) or None


def _initials(raw):
    if not isinstance(raw, str):
        return None
    s = re.sub(r"[’'`]", "", raw.upper())
    ini = "".join(t for t in s.split() if len(t) == 1)
    return ini or None


def build_dataframe(records):
    df = pd.DataFrame.from_records(records)
    df["sequence"] = pd.to_numeric(df["sequence"], errors="coerce").astype("Int16")
    df["flight_date"] = pd.to_datetime(df["flight_date"], format="%Y-%m-%d", errors="coerce")
    df["departure_dt"] = pd.to_datetime(
        df["flight_date"].dt.strftime("%Y-%m-%d") + " " + df["flight_time"],
        format="%Y-%m-%d %H:%M",
        errors="coerce",
    )
    df["gender"] = df["title"].str.upper().map(GENDER)
    ff = df["ff_raw"].str.extract(r"^\s*([A-Z0-9]{2})\s*[-\s]\s*(\d+)\s*$")
    df["ff_program"] = ff[0]
    df["ff_number"] = ff[1]
    df["has_ff"] = df["ff_raw"].notna()

    df["name_key"] = df["passenger_name_raw"].map(_name_key)
    df["name_initials"] = df["passenger_name_raw"].map(_initials)

    df["route"] = df["from_iata"].fillna("?") + "-" + df["to_iata"].fillna("?")
    df["pnr_group_size"] = df.groupby("pnr")["pnr"].transform("size").astype("Int16")

    for c in ["title", "gender", "travel_class", "ff_program", "from_iata", "to_iata",
              "flight_no", "route", "source_file"]:
        df[c] = df[c].astype("category")

    cols = [
        "source_file", "sheet", "sheet_idx",
        "flight_no", "flight_date", "flight_time", "departure_dt",
        "from_city", "from_iata", "to_city", "to_iata", "route",
        "passenger_name_raw", "name_key", "name_initials", "title", "gender",
        "travel_class", "ff_raw", "ff_program", "ff_number", "has_ff",
        "pnr", "pnr_group_size", "ticket_no", "sequence",
        "gate", "seat", "operator_note",
    ]
    return df[cols]


def drop_constant(df):
    dropped = []
    for c in CONSTANT_CANDIDATES:
        if c in df and df[c].nunique(dropna=True) <= 1:
            dropped.append((c, df[c].dropna().unique().tolist()))
    df = df.drop(columns=[c for c, _ in dropped])
    return df, dropped


def quality_report(df):
    print("\n=== Отчёт о качестве данных ===")
    print(f"строк (талонов): {len(df):,}   файлов: {df['source_file'].nunique()}")
    print(f"период: {df['flight_date'].min():%Y-%m-%d} .. {df['flight_date'].max():%Y-%m-%d}")

    nulls = df.isna().mean().round(3)
    nulls = nulls[nulls > 0]
    if len(nulls):
        print("доля пропусков:", ", ".join(f"{k}={v:.1%}" for k, v in nulls.items()))

    dup_t = df["ticket_no"].duplicated(keep=False).sum()
    dup_p = df.duplicated(subset=["pnr", "ticket_no", "flight_no", "flight_date"], keep=False).sum()
    print(f"дубликаты номера билета: {dup_t}, полные дубликаты талонов: {dup_p}")
    air = pd.concat([
        df[["from_iata", "from_city"]].set_axis(["iata", "city"], axis=1),
        df[["to_iata", "to_city"]].set_axis(["iata", "city"], axis=1),
    ]).astype(str)
    n = air.groupby("iata")["city"].nunique()
    bad = n[n > 1]
    if len(bad):
        print("IATA-коды с несколькими городами:")
        for code in bad.index:
            print(f"  {code}: {sorted(air.loc[air.iata == code, 'city'].unique())}")

    multi = (df.dropna(subset=["name_key"])
               .groupby(["name_key", "flight_date"], observed=True)["flight_no"].nunique())
    print(f"пар (имя, день) с >1 рейсом: {(multi > 1).sum()}  (могут быть однофамильцы)")


def collect_files(root: Path):
    files = []
    for p in root.rglob("*"):
        if not p.is_file() or p.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        if p.name.startswith(("~$", "._")) or "__MACOSX" in p.parts:
            continue
        files.append(p)
    return sorted(files)


def _parse_all(files, workers):
    paths = [str(p) for p in files]
    workers = min(max(1, int(workers)), len(paths))
    records, errors = [], []
    if workers == 1:
        for recs, errs in map(parse_workbook, paths):
            records += recs
            errors += errs
    else:
        with ProcessPoolExecutor(max_workers=workers) as ex:
            for recs, errs in ex.map(parse_workbook, paths, chunksize=1):
                records += recs
                errors += errs
    return records, errors


def parse_boardingPass_zip(zip_path, output_path="boarding_passes.parquet", *,
                           workers=4, keep_constant=False, verbose=True):
    log = print if verbose else (lambda *a, **k: None)

    src = Path(zip_path)
    if not src.exists():
        raise FileNotFoundError(f"Не найден архив или папка: {src}")

    out = Path(output_path)
    if out.suffix == "":
        out = out.with_suffix(".parquet")
    out.parent.mkdir(parents=True, exist_ok=True)


    with tempfile.TemporaryDirectory() as tmp:
        if src.is_file():
            if src.suffix.lower() != ".zip":
                raise ValueError(f"Ожидался .zip-архив или папка, получено: {src.name}")
            with zipfile.ZipFile(src) as z:
                z.extractall(tmp)
            root = Path(tmp)
        else:
            root = src

        files = collect_files(root)
        if not files:
            raise ValueError(f"В {src} не найдено ни одной книги .xlsx/.xlsm")
        log(f"Найдено книг: {len(files)}")

        records, errors = _parse_all(files, workers)

    if errors:
        err_path = out.with_name(out.stem + "_parse_errors.csv")
        pd.DataFrame(errors).to_csv(err_path, index=False)
        log(f"Проблемных листов: {len(errors)} -> {err_path}")

    if not records:
        raise ValueError("Не удалось распарсить ни одного талона (см. файл с ошибками)")

    df = build_dataframe(records)

    if not keep_constant:
        df, dropped = drop_constant(df)
        for c, vals in dropped:
            log(f"колонка '{c}' константна ({vals}) -> удалена")

    if verbose:
        quality_report(df)

    df.to_parquet(out, index=False, compression="zstd")
    log(f"\nГотово: {out}  ({len(df):,} строк, {df.shape[1]} колонок)")
    return df
